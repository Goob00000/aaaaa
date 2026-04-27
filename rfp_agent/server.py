"""
RFP Response Agent — FastAPI web server.

Run from the rfp_agent/ directory:
    uvicorn server:app --reload --port 8000
"""

import asyncio
import io
import os
import sys
import threading
from pathlib import Path

import chromadb
import pandas as pd
import uvicorn
from chromadb.utils.embedding_functions import SentenceTransformerEmbeddingFunction
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

BASE_DIR = Path(__file__).parent
sys.path.insert(0, str(BASE_DIR))

from agent import RFPAgent
from database import (
    create_job, finish_job, get_job, get_job_categories, get_job_stats,
    get_library_categories, get_library_count, get_library_entries,
    get_results, init_db, insert_result, list_jobs, update_job_progress,
    update_result, upsert_library_entries,
)
from ingest import ingest, load_library

# ── App setup ─────────────────────────────────────────────────────────────────

app = FastAPI(title="RFP Response Agent")
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))

UPLOAD_DIR = BASE_DIR / "data" / "uploads"
CHROMA_PATH = str(BASE_DIR / "knowledge_base")
REVIEW_STATUSES = ["Draft", "Pending Review", "Approved", "Published"]


@app.on_event("startup")
async def startup():
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    init_db()


# ── Helpers ───────────────────────────────────────────────────────────────────

def _library_loaded() -> bool:
    try:
        embed_fn = SentenceTransformerEmbeddingFunction(model_name="all-MiniLM-L6-v2")
        client = chromadb.PersistentClient(path=CHROMA_PATH)
        col = client.get_collection(name="rfp_predefined_responses", embedding_function=embed_fn)
        return col.count() > 0
    except Exception:
        return False


def _process_job_thread(job_id: str, questions: list[str], updated_by: str):
    try:
        agent = RFPAgent(chroma_path=CHROMA_PATH)
        for i, question in enumerate(questions):
            result = agent.process_question(question, updated_by=updated_by)
            insert_result(job_id, result)
            update_job_progress(job_id, processed=i + 1)
        finish_job(job_id)
    except Exception as exc:
        finish_job(job_id, error=str(exc))


def _build_excel(results: list[dict]) -> io.BytesIO:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    cols = ["original_question", "detected_category", "matched_topic",
            "draft_answer", "updated_by", "review_status", "match_score", "needs_review"]
    df = pd.DataFrame(results)[[c for c in cols if c in pd.DataFrame(results).columns]]

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="RFP Responses")
        ws = writer.sheets["RFP Responses"]

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

        ws.freeze_panes = "A2"

        widths = {"original_question": 50, "detected_category": 22, "matched_topic": 30,
                  "draft_answer": 80, "updated_by": 18, "review_status": 18,
                  "match_score": 14, "needs_review": 14}
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 15)

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row[0].row].height = 60

        if "review_status" in df.columns:
            ci = list(df.columns).index("review_status") + 1
            cl = get_column_letter(ci)
            dv = DataValidation(type="list", formula1='"Draft,Pending Review,Approved,Published"', showDropDown=False)
            dv.sqref = f"{cl}2:{cl}{len(df) + 1}"
            ws.add_data_validation(dv)

        yellow = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
        for i, row in enumerate(results, start=2):
            if row.get("needs_review"):
                for ci in range(1, len(df.columns) + 1):
                    ws.cell(row=i, column=ci).fill = yellow

    buf.seek(0)
    return buf


# ── Home ──────────────────────────────────────────────────────────────────────

@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    return templates.TemplateResponse("index.html", {
        "request": request,
        "jobs": list_jobs(),
        "library_loaded": _library_loaded(),
        "library_count": get_library_count(),
        "api_key_set": bool(os.environ.get("ANTHROPIC_API_KEY")),
    })


# ── Library ───────────────────────────────────────────────────────────────────

@app.post("/ingest")
async def ingest_library(
    file: UploadFile = File(...),
    reset: str = Form(default="true"),
):
    if not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "File must be a .xlsx spreadsheet")

    content = await file.read()
    tmp_path = UPLOAD_DIR / file.filename
    tmp_path.write_bytes(content)

    do_reset = reset.lower() != "false"
    categorized_df = await asyncio.to_thread(ingest, str(tmp_path), do_reset)
    entries = categorized_df[["topic", "response", "category", "keywords"]].to_dict("records")
    upsert_library_entries(entries)

    return RedirectResponse("/library?ingested=1", status_code=303)


@app.get("/library", response_class=HTMLResponse)
async def library_page(
    request: Request,
    category: str = "",
    search: str = "",
    page: int = 1,
):
    limit = 50
    offset = (page - 1) * limit
    entries, total = get_library_entries(
        category=category or None,
        search=search or None,
        limit=limit,
        offset=offset,
    )
    total_pages = max(1, (total + limit - 1) // limit)
    return templates.TemplateResponse("library.html", {
        "request": request,
        "entries": entries,
        "categories": get_library_categories(),
        "selected_category": category,
        "search": search,
        "page": page,
        "total_pages": total_pages,
        "total": total,
        "ingested": request.query_params.get("ingested") == "1",
    })


# ── RFP Processing ────────────────────────────────────────────────────────────

@app.post("/process")
async def process_rfp(
    file: UploadFile = File(...),
    updated_by: str = Form(default="AI Draft"),
):
    if not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "File must be a .xlsx spreadsheet")

    content = await file.read()
    df = pd.read_excel(io.BytesIO(content))
    df.columns = [c.strip().lower().replace(" ", "_") for c in df.columns]

    if "question" not in df.columns:
        raise HTTPException(400, "Spreadsheet must have a 'question' column")

    questions = (
        df["question"].dropna().astype(str).str.strip().pipe(lambda s: s[s != ""])
    ).tolist()

    if not questions:
        raise HTTPException(400, "No questions found in the spreadsheet")

    (UPLOAD_DIR / file.filename).write_bytes(content)
    job_id = create_job(file.filename, total=len(questions))

    threading.Thread(
        target=_process_job_thread,
        args=(job_id, questions, updated_by),
        daemon=True,
    ).start()

    return RedirectResponse(f"/jobs/{job_id}", status_code=303)


# ── Job views ─────────────────────────────────────────────────────────────────

@app.get("/jobs/{job_id}", response_class=HTMLResponse)
async def job_page(
    request: Request,
    job_id: str,
    category: str = "",
    status: str = "",
):
    job = get_job(job_id)
    if not job:
        raise HTTPException(404, "Job not found")

    results = get_results(job_id, category=category or None, status=status or None)
    stats = get_job_stats(job_id)

    return templates.TemplateResponse("job.html", {
        "request": request,
        "job": job,
        "results": results,
        "stats": stats,
        "categories": get_job_categories(job_id),
        "review_statuses": REVIEW_STATUSES,
        "selected_category": category,
        "selected_status": status,
    })


@app.get("/jobs/{job_id}/status")
async def job_status(job_id: str):
    job = get_job(job_id)
    if not job:
        raise HTTPException(404)
    return {
        "status": job["status"],
        "processed": job["processed"],
        "total": job["total"],
        "error": job["error_message"],
    }


@app.patch("/results/{result_id}")
async def patch_result(result_id: int, request: Request):
    body = await request.json()
    if not update_result(result_id, **body):
        raise HTTPException(400, "No valid fields provided")
    return {"ok": True}


@app.get("/jobs/{job_id}/download")
async def download_job(job_id: str):
    job = get_job(job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    results = get_results(job_id)
    if not results:
        raise HTTPException(404, "No results available")

    buf = await asyncio.to_thread(_build_excel, results)
    safe_name = f"rfp_response_{Path(job['filename']).stem}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}"'},
    )


# ── Entry point ───────────────────────────────────────────────────────────────

if __name__ == "__main__":
    uvicorn.run("server:app", host="0.0.0.0", port=8000, reload=True)
